"""Build outputs/asset_pnl_variance.xlsx: an asset-management view of the same wind farm.

Budget (P50 proxy) is set ex ante from 2014 data and a 1999 to 2013 ERA5 wind index; 2015 is then
analysed as the operating year: monthly budget vs actual, a variance bridge by driver, a turbine
ranking, and a reforecast. Every number in the analysis sheets is an Excel formula.

Run from the repo root after the pipeline:  python -m src.wpa.build_asset_pnl
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "asset_pnl_variance.xlsx"

F = "Arial"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
YELLOW = PatternFill("solid", start_color="FFFF00")
HEAD = PatternFill("solid", start_color="1D2733")
SUB = PatternFill("solid", start_color="E8ECEF")
THIN = Border(bottom=Side(style="thin", color="B0B8BF"))
TOP = Border(top=Side(style="thin", color="1D2733"))
MWH = '#,##0.0;(#,##0.0);"-"'
EUR = '€#,##0;(€#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=F, color=color, bold=bold, size=size, italic=italic)


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(name=F, color="FFFFFF", bold=True, size=10)
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = sub
    ws["A2"].font = font(color="5B6B7A", italic=True)


def wind_index() -> pd.DataFrame:
    """Monthly mean of a capped cubic wind energy proxy from ERA5 100 m wind speed."""
    e = pd.read_csv(ROOT / "data" / "era5_wind_la_haute_borne.csv", usecols=["datetime", "ws_100m"],
                    parse_dates=["datetime"])
    e["p"] = np.clip((e["ws_100m"] - 3).clip(lower=0) ** 3 / 9 ** 3, 0, 1)
    e["y"], e["m"] = e["datetime"].dt.year, e["datetime"].dt.month
    lt = e[e["y"].between(1999, 2013)].groupby("m")["p"].mean()
    return pd.DataFrame({"m": range(1, 13), "lt": lt.values,
                         "y2014": e[e["y"] == 2014].groupby("m")["p"].mean().values,
                         "y2015": e[e["y"] == 2015].groupby("m")["p"].mean().values})


def main():
    lm = pd.read_csv(ROOT / "outputs" / "loss_monthly.csv", parse_dates=["month"])
    wi = wind_index()
    turbines = sorted(lm["turbine"].unique())
    wb = Workbook()

    # ---------------- README ----------------
    ws = wb.active
    ws.title = "README"
    title(ws, "La Haute Borne wind farm: asset P&L and variance analysis, 2015",
          "Companion to github.com/syulimo/wind-performance-analysis")
    rows = [
        ("Purpose", "Asset-management view of the same SCADA analysis: budget vs actual, variance bridge by driver, turbine ranking, reforecast."),
        ("Operating year", "2015. Budget is set ex ante from 2014 data only, so 2015 is analysed out of sample."),
        ("Budget (P50 proxy)", "2014 potential energy, long-term corrected with an ERA5 wind index (1999 to 2013), spread by the long-term monthly profile."),
        ("Revenue", "Energy x tariff. The tariff on Inputs is an editable assumption; change it to see every euro figure update."),
        ("Data sources", "outputs/loss_monthly.csv from the pipeline (ENGIE La Haute Borne SCADA via NREL OpenOA); ERA5 100 m wind speed from the same OpenOA dataset."),
        ("Not included", "O&M, land, insurance and other costs: no public cost data exist for this site, so the P&L stops at revenue and lost revenue."),
        ("Colour legend", "Blue text = hardcoded input or imported data. Black = formula. Green = direct link to another sheet. Yellow fill = assumption to edit."),
        ("How to use", "Edit the yellow cells on Inputs (tariff, months of actuals). Everything else recalculates."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=k).font = font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110

    # ---------------- Inputs ----------------
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs and budget", "Edit yellow cells only")
    header(ws, 4, ["Item", "Value", "Unit", "Source / note"])
    ws["A5"], ws["B5"], ws["C5"] = "Tariff", 82, "€/MWh"
    ws["D5"] = ("Assumption: French onshore wind feed-in tariff base level (arrêté of 17 Nov 2008, 8.2 c€/kWh), "
                "before annual indexation. Edit to test other prices or a PPA.")
    ws["A6"], ws["B6"], ws["C6"] = "Months of actuals for reforecast", 6, "months"
    ws["D6"] = "Reforecast sheet uses actuals for months 1 to this number, budget after. Enter 1 to 11."
    for r in (5, 6):
        ws[f"B{r}"].font = font(BLUE)
        ws[f"B{r}"].fill = YELLOW
    ws["A8"] = "Budget build (formulas)"
    ws["A8"].font = font(bold=True)
    ws["A9"], ws["C9"] = "2014 potential energy, fleet", "MWh"
    ws["B9"] = '=SUMIFS(Data_Monthly!$E:$E,Data_Monthly!$I:$I,2014)'
    ws["D9"] = "Expected energy from the reference power curves, 2014 (pipeline output)"
    ws["A10"], ws["C10"] = "Long-term / 2014 wind index ratio", "x"
    ws["B10"] = '=SUM(Wind_Index!$B$5:$B$16)/SUM(Wind_Index!$C$5:$C$16)'
    ws["D10"] = "ERA5 1999 to 2013 mean over 2014; above 1 means 2014 was a low-wind year"
    ws["A11"], ws["C11"] = "Annual budget, P50 proxy", "MWh"
    ws["B11"] = "=B9*B10"
    ws["D11"] = "Spread across months by the long-term monthly profile on Wind_Index"
    ws["A12"], ws["C12"] = "Annual budget revenue", "€"
    ws["B12"] = "=B11*B5"
    for r, fmt in ((9, MWH), (10, "0.000x"), (11, MWH), (12, EUR)):
        ws[f"B{r}"].number_format = fmt
        ws[f"B{r}"].font = font()
    ws["B5"].number_format = '€#,##0.00'
    for col, w in zip("ABCD", (34, 14, 10, 100)):
        ws.column_dimensions[col].width = w
    for r in range(5, 13):
        for col in "ACD":
            ws[f"{col}{r}"].font = font(bold=(col == "A" and r in (11, 12)))

    # ---------------- Wind_Index ----------------
    ws = wb.create_sheet("Wind_Index")
    title(ws, "ERA5 wind energy index by month",
          "Mean of min(max(ws100 - 3, 0)^3 / 9^3, 1), hourly ERA5 100 m wind speed at the site; computed in src/wpa/build_asset_pnl.py")
    header(ws, 4, ["Month", "Long-term 1999-2013", "2014", "2015", "Budget share of year", "2015 vs long-term"])
    for i, r in wi.iterrows():
        row = 5 + i
        ws.cell(row=row, column=1, value=MONTHS[i]).font = font()
        for j, key in enumerate(("lt", "y2014", "y2015"), start=2):
            c = ws.cell(row=row, column=j, value=round(float(r[key]), 6))
            c.font = font(BLUE)
            c.number_format = "0.000"
        ws.cell(row=row, column=5, value=f"=B{row}/SUM($B$5:$B$16)").number_format = PCT
        ws.cell(row=row, column=6, value=f"=D{row}/B{row}-1").number_format = PCT
        for col in (5, 6):
            ws.cell(row=row, column=col).font = font()
    ws["A17"], ws["B17"], ws["C17"], ws["D17"] = "Year", "=SUM(B5:B16)", "=SUM(C5:C16)", "=SUM(D5:D16)"
    ws["F17"] = "=D17/B17-1"
    for col in "ABCDF":
        ws[f"{col}17"].font = font(bold=True)
        ws[f"{col}17"].border = TOP
    for col in "BCD":
        ws[f"{col}17"].number_format = "0.000"
    ws["F17"].number_format = PCT
    for col, w in zip("ABCDEF", (10, 20, 12, 12, 20, 18)):
        ws.column_dimensions[col].width = w

    # ---------------- Data_Monthly ----------------
    ws = wb.create_sheet("Data_Monthly")
    title(ws, "Monthly energy by turbine (pipeline output)",
          "Source: outputs/loss_monthly.csv, written by src/wpa/pipeline.py from 10-minute SCADA. MWh.")
    cols = ["Turbine", "Month", "Actual", "", "Potential", "Downtime", "Underperformance", "Icing", "Year",
            "Month no.", "Curtailment"]
    header(ws, 4, ["Turbine", "Month", "Actual MWh", "(blank)", "Potential MWh", "Downtime MWh",
                   "Underperf. MWh", "Icing MWh", "Year", "Month no.", "Curtailment MWh"])
    for i, r in lm.reset_index(drop=True).iterrows():
        row = 5 + i
        ws.cell(row=row, column=1, value=r["turbine"])
        ws.cell(row=row, column=2, value=r["month"].to_pydatetime()).number_format = "yyyy-mm"
        vals = {3: r["actual_mwh"], 5: r["potential_mwh"], 6: r["downtime_mwh"], 7: r["underperf_mwh"],
                8: r["icing_mwh"], 11: r["curtailment_mwh"]}
        for col, v in vals.items():
            c = ws.cell(row=row, column=col, value=round(float(v), 3))
            c.number_format = MWH
        ws.cell(row=row, column=9, value=f"=YEAR(B{row})")
        ws.cell(row=row, column=10, value=f"=MONTH(B{row})")
        for col in range(1, 12):
            c = ws.cell(row=row, column=col)
            c.font = font(BLACK if col in (9, 10) else BLUE)
    last_data = 4 + len(lm)
    for col, w in zip("ABCDEFGHIJK", (10, 10, 12, 8, 13, 13, 14, 11, 7, 9, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    def s(col, year, month_ref=None, turbine_ref=None):
        crit = f"Data_Monthly!${col}$5:${col}${last_data},Data_Monthly!$I$5:$I${last_data},{year}"
        if month_ref:
            crit += f",Data_Monthly!$J$5:$J${last_data},{month_ref}"
        if turbine_ref:
            crit += f",Data_Monthly!$A$5:$A${last_data},{turbine_ref}"
        return f"SUMIFS({crit})"

    # ---------------- Fleet_2015 ----------------
    ws = wb.create_sheet("Fleet_2015")
    title(ws, "Fleet budget vs actual by month, 2015",
          "Losses shown negative. Check column must be zero: resource + losses + other = variance to budget.")
    header(ws, 4, ["Month no.", "Month", "Budget MWh", "Potential MWh", "Wind resource var.", "Downtime",
                   "Underperf.", "Icing", "Curtailment", "Actual MWh", "Other", "Variance to budget",
                   "Check", "Budget revenue", "Actual revenue", "Revenue variance"])
    for m in range(1, 13):
        r = 4 + m
        ws.cell(row=r, column=1, value=m).font = font(BLUE)
        ws.cell(row=r, column=2, value=MONTHS[m - 1]).font = font()
        f = {
            3: f"=Inputs!$B$11*Wind_Index!E{r}",
            4: "=" + s("E", 2015, f"$A{r}"),
            5: f"=D{r}-C{r}",
            6: "=-" + s("F", 2015, f"$A{r}"),
            7: "=-" + s("G", 2015, f"$A{r}"),
            8: "=-" + s("H", 2015, f"$A{r}"),
            9: "=-" + s("K", 2015, f"$A{r}"),
            10: "=" + s("C", 2015, f"$A{r}"),
            11: f"=J{r}-D{r}-SUM(F{r}:I{r})",
            12: f"=J{r}-C{r}",
            13: f"=ROUND(E{r}+SUM(F{r}:I{r})+K{r}-L{r},6)",
            14: f"=C{r}*Inputs!$B$5",
            15: f"=J{r}*Inputs!$B$5",
            16: f"=O{r}-N{r}",
        }
        for col, formula in f.items():
            c = ws.cell(row=r, column=col, value=formula)
            c.font = font()
            c.number_format = EUR if col >= 14 else MWH
    ws["B17"] = "2015"
    for col in range(3, 17):
        L = get_column_letter(col)
        ws[f"{L}17"] = f"=SUM({L}5:{L}16)"
        ws[f"{L}17"].number_format = EUR if col >= 14 else MWH
    for col in range(1, 17):
        c = ws.cell(row=17, column=col)
        c.font = font(bold=True)
        c.border = TOP
    ws["A19"] = "Variance to budget, %"
    ws["C19"] = "=IFERROR(L17/C17,0)"
    ws["C19"].number_format = PCT
    ws["A19"].font = ws["C19"].font = font(bold=True)
    ws["A20"] = "Reading: the wind resource column is what the weather did; downtime to curtailment are what operations lost; other is curve scatter and records excluded by data quality."
    ws["A20"].font = font(color="5B6B7A", italic=True)
    for i, w in enumerate((9, 8, 12, 13, 13, 11, 11, 9, 11, 12, 10, 13, 8, 14, 14, 14), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C5"

    # ---------------- Bridge ----------------
    ws = wb.create_sheet("Bridge")
    title(ws, "2015 revenue bridge: budget to actual", "Each step links to Fleet_2015 totals; euros = MWh x tariff")
    header(ws, 4, ["Step", "MWh", "€", "Share of budget"])
    steps = [("Budget", "=Fleet_2015!C17"), ("Wind resource", "=Fleet_2015!E17"), ("Downtime", "=Fleet_2015!F17"),
             ("Underperformance", "=Fleet_2015!G17"), ("Icing", "=Fleet_2015!H17"),
             ("Curtailment", "=Fleet_2015!I17"), ("Other", "=Fleet_2015!K17"), ("Actual", "=Fleet_2015!J17")]
    for i, (lab, ref) in enumerate(steps):
        r = 5 + i
        ws.cell(row=r, column=1, value=lab).font = font(bold=lab in ("Budget", "Actual"))
        c = ws.cell(row=r, column=2, value=ref)
        c.font, c.number_format = font(GREEN), MWH
        c = ws.cell(row=r, column=3, value=f"=B{r}*Inputs!$B$5")
        c.font, c.number_format = font(), EUR
        c = ws.cell(row=r, column=4, value=f"=IFERROR(B{r}/$B$5,0)")
        c.font, c.number_format = font(), PCT
    ws["A14"] = "Check: budget + steps - actual"
    ws["B14"] = "=ROUND(B5+SUM(B6:B11)-B12,6)"
    ws["B14"].number_format = MWH
    ws["A15"] = "Operational losses, € (downtime to curtailment)"
    ws["C15"] = "=SUM(C7:C10)"
    ws["C15"].number_format = EUR
    ws["A16"] = "Largest operational driver"
    ws["C16"] = '=INDEX(A7:A10,MATCH(MIN(C7:C10),C7:C10,0))'
    for r in (14, 15, 16):
        ws[f"A{r}"].font = font(bold=True)
        for col in "BC":
            ws[f"{col}{r}"].font = font()
    chart = BarChart()
    chart.type = "col"
    chart.title = "2015 variance to budget by driver (€)"
    chart.y_axis.title = "€"
    chart.y_axis.numFmt = '€#,##0'
    chart.add_data(Reference(ws, min_col=3, min_row=6, max_row=11), titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=11))
    chart.legend = None
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "F4")
    for col, w in zip("ABCD", (44, 12, 14, 16)):
        ws.column_dimensions[col].width = w

    # ---------------- Turbine_2015 ----------------
    ws = wb.create_sheet("Turbine_2015")
    title(ws, "Turbine ranking by lost energy and lost revenue, 2015", "Rank 1 = most lost energy; act on it first")
    header(ws, 4, ["Turbine", "Actual MWh", "Potential MWh", "Downtime", "Underperf.", "Icing", "Curtailment",
                   "Total lost MWh", "Lost revenue", "Share of fleet loss", "Rank", "Underperf. 2014",
                   "Underperf. change"])
    n = len(turbines)
    for i, t in enumerate(turbines):
        r = 5 + i
        ws.cell(row=r, column=1, value=t).font = font(BLUE)
        f = {2: "=" + s("C", 2015, turbine_ref=f"$A{r}"), 3: "=" + s("E", 2015, turbine_ref=f"$A{r}"),
             4: "=" + s("F", 2015, turbine_ref=f"$A{r}"), 5: "=" + s("G", 2015, turbine_ref=f"$A{r}"),
             6: "=" + s("H", 2015, turbine_ref=f"$A{r}"), 7: "=" + s("K", 2015, turbine_ref=f"$A{r}"),
             8: f"=SUM(D{r}:G{r})", 9: f"=H{r}*Inputs!$B$5", 10: f"=IFERROR(H{r}/SUM($H$5:$H${4 + n}),0)",
             11: f"=RANK(H{r},$H$5:$H${4 + n},0)", 12: "=" + s("G", 2014, turbine_ref=f"$A{r}"),
             13: f"=E{r}-L{r}"}
        for col, formula in f.items():
            c = ws.cell(row=r, column=col, value=formula)
            c.font = font()
            c.number_format = {9: EUR, 10: PCT, 11: "0"}.get(col, MWH)
    tr = 5 + n
    ws.cell(row=tr, column=1, value="Fleet")
    for col in (2, 3, 4, 5, 6, 7, 8, 9, 12, 13):
        L = get_column_letter(col)
        ws.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr - 1})").number_format = EUR if col == 9 else MWH
    for col in range(1, 14):
        c = ws.cell(row=tr, column=col)
        c.font = font(bold=True)
        c.border = TOP
    for i, w in enumerate((10, 12, 13, 11, 11, 9, 11, 13, 13, 14, 7, 14, 15), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---------------- Reforecast ----------------
    ws = wb.create_sheet("Reforecast")
    title(ws, "Full-year 2015 reforecast", "Actuals through the month set on Inputs!B6, budget for the rest of the year")
    header(ws, 4, ["Line", "MWh", "€", "Note"])
    lines = [
        ("Months of actuals", "=Inputs!B6", None, "Set on Inputs"),
        ("YTD budget", "=SUMIFS(Fleet_2015!C5:C16,Fleet_2015!A5:A16,\"<=\"&B5)", True, ""),
        ("YTD actual", "=SUMIFS(Fleet_2015!J5:J16,Fleet_2015!A5:A16,\"<=\"&B5)", True, ""),
        ("YTD potential", "=SUMIFS(Fleet_2015!D5:D16,Fleet_2015!A5:A16,\"<=\"&B5)", True, "Energy the turbines could have made in the actual wind"),
        ("YTD operating ratio", "=IFERROR(B7/B8,0)", None, "Actual / potential: share of available wind energy converted"),
        ("Rest-of-year budget", "=SUMIFS(Fleet_2015!C5:C16,Fleet_2015!A5:A16,\">\"&B5)", True, ""),
        ("Reforecast A: YTD actual + budget", "=B7+B10", True, "Assumes rest of year at budget"),
        ("Reforecast B: YTD actual + budget x operating ratio", "=B7+B10*B9", True, "Carries YTD operational losses forward; normal wind assumed"),
        ("Full-year budget", "=Fleet_2015!C17", True, ""),
        ("Full-year actual (for back-test)", "=Fleet_2015!J17", True, "Known here because 2015 is historical"),
        ("Error, reforecast A vs actual", "=B11-B14", True, ""),
        ("Error, reforecast B vs actual", "=B12-B14", True, ""),
    ]
    for i, (lab, formula, money, note) in enumerate(lines):
        r = 5 + i
        ws.cell(row=r, column=1, value=lab).font = font(bold=lab.startswith("Reforecast"))
        c = ws.cell(row=r, column=2, value=formula)
        c.font = font(GREEN if formula.startswith("=Inputs") or formula.startswith("=Fleet_2015!") else BLACK)
        c.number_format = "0" if i == 0 else ("0.0%" if lab == "YTD operating ratio" else MWH)
        if money:
            c2 = ws.cell(row=r, column=3, value=f"=B{r}*Inputs!$B$5")
            c2.font, c2.number_format = font(), EUR
        ws.cell(row=r, column=4, value=note).font = font(color="5B6B7A", italic=True)
    for col, w in zip("ABCD", (48, 12, 14, 64)):
        ws.column_dimensions[col].width = w

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    wb["Inputs"]["B5"].comment = Comment("Assumption, not site revenue data. Edit freely.", "model")
    wb.active = wb.sheetnames.index("Bridge")
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
